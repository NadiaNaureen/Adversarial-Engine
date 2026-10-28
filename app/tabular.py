"""Safe, in-memory ingestion of user-provided CSV and Excel tables."""
import csv
import io
from pathlib import Path

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from .models import ArsenalRecord, TargetRecord

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
TARGET_COLUMNS = {"id", "age", "zip_code", "occupation", "gender", "medical_flag"}
ARSENAL_COLUMNS = TARGET_COLUMNS | {"name"}


def _clean_row(row: dict) -> dict:
    return {str(key).strip(): "" if value is None else str(value).strip() for key, value in row.items() if key is not None}


def _xlsx_rows(content: bytes) -> list[dict]:
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []
    return [_clean_row(dict(zip(headers, row))) for row in rows if any(value is not None for value in row)]


async def parse_upload(file: UploadFile, record_type: type[BaseModel]) -> list[BaseModel]:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xlsm"}:
        raise HTTPException(400, "Upload a .csv, .xlsx, or .xlsm file.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File is larger than the 10 MB upload limit.")
    try:
        rows = ([_clean_row(row) for row in csv.DictReader(io.StringIO(content.decode("utf-8-sig")))] if suffix == ".csv" else _xlsx_rows(content))
    except Exception as exc:
        raise HTTPException(400, f"Could not read {file.filename}: {exc}") from exc
    required = ARSENAL_COLUMNS if record_type is ArsenalRecord else TARGET_COLUMNS
    missing = sorted(required - (set(rows[0]) if rows else set()))
    if missing:
        raise HTTPException(400, f"Missing required column(s): {', '.join(missing)}")
    try:
        return [record_type.model_validate(row) for row in rows]
    except ValidationError as exc:
        first = exc.errors()[0]
        raise HTTPException(400, f"Invalid row data: {first['loc'][0]} - {first['msg']}") from exc
