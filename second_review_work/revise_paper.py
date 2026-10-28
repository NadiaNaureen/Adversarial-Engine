from pathlib import Path
from collections import defaultdict
from copy import deepcopy
import csv, json, statistics, math
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
RESULTS=ROOT/'results'
SOURCE=Path(r'C:\Users\DELL\Downloads\AI-Driven Privacy Risk Assessment and Record Linkage Auditing Using Large Language Models.docx')
OUTPUT=ROOT/'LLM_Linkage_Audit_Revised_6_Pages.docx'
SEEDS=[11,22,33,44,55]
POLICIES=['original','age_band','postal_prefix','medical_suppression','occupation_suppression','combined','k5']
NAMES=['Original','Age bands','Postal prefix','Medical removed','Occupation removed','Combined','Global k≥5']
def read(seed,file):
    with (RESULTS/str(seed)/(file+'.csv')).open(encoding='utf-8') as f:return list(csv.DictReader(f))
def metric(rows):
    c=w=pa=f=aa=0
    for r in rows:
        truth=r['truth'];pred=r['prediction']
        if truth:
            if pred==truth:c+=1
            elif pred:w+=1
            else:pa+=1
        elif pred:f+=1
        else:aa+=1
    assert c+w+pa+f+aa==len(rows)
    return dict(C=c,W=w,P_A=pa,F=f,A_A=aa,recall=c/(c+w+pa),fpr=f/(f+aa),precision=c/(c+w+f) if c+w+f else None)
def stats(file,method,condition=None):
    rr=[];mm=[]
    for seed in SEEDS:
        rows=[r for r in read(seed,file) if r['method']==method and (condition is None or r['condition']==condition)]
        mm.append(metric(rows));rr+=rows
    return dict(totals=metric(rr),**{k:[m[k] for m in mm] for k in ['recall','fpr','precision']})
def ms(vals):
    vals=[v for v in vals if v is not None]
    return '—' if not vals else f'{100*statistics.mean(vals):.1f} ± {100*statistics.stdev(vals):.1f}'
def mean(vals):return 100*statistics.mean(vals)
def fmt(vals):return f'{mean(vals):.1f}'
MAIN={m:stats('predictions',m) for m in ['exact','fs','local','no_margin','top1']}
MIT={p:stats('mitigation_predictions','local',p) for p in POLICIES}
ADAPT={p:{m:stats('adaptive_predictions',m,p) for m in ['adaptive_local','adaptive_fs']} for p in POLICIES}
UTILITY={p:[next(r for r in read(s,'utility') if r['policy']==p) for s in SEEDS] for p in POLICIES}
MATCH={}
for m in ['local','fs']:
    for budget in [.02,.05]:
        rows=[r for s in SEEDS for r in read(s,'matched_fpr') if r['method']==m and float(r['budget'])==budget]
        MATCH[m,budget]=[float(r['recall']) for r in rows]

doc=Document(SOURCE)
ps=list(doc.paragraphs)
def replace(i,text):
    p=ps[i]
    # Retain paragraph and first-run properties, including the conference template.
    first=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();run=p.add_run(text)
    if first is not None:run._r.insert(0,first)
def cell(c,text):
    p=c.paragraphs[0]
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(str(text))
    if props is not None:r._r.insert(0,props)

replace(0,'LLM-Assisted Record-Linkage Auditing with Verified Evidence, Abstention, and Mitigation Re-Audits')
local=MAIN['local'];fs=MAIN['fs'];counts=local['totals']
replace(2,f'Abstract—Removing names and other direct identifiers does not prevent records from being linked through demographic details. We present an LLM-assisted audit engine that retrieves candidates, asks Gemini to assess them, checks the field evidence, and either links a record or withholds a decision. The same index supports repeat audits after generalization and suppression. We evaluate five synthetic populations of 10,000 identities, with separate validation and test identities and six target conditions. Across 18,000 test variants, local retrieval and verification achieve {ms(local["recall"])}% correct-link recall and {ms(local["fpr"])}% false links for absent targets. We compare recall at matched false-link limits and test attackers that adjust to each mitigation. In the unchanged cloud subset, Gemini and its local control both reach 43.3% recall, with false-link rates of 1.7% and 5.0%. A small semantic pilot shows no retrieval advantage on these structured profiles. The findings support an auditable workflow, while showing the limits of literal field checks and fixed attacker settings.')
replace(19,'Figure 1 shows the application. (1) The web interface and FastAPI backend accept uploaded or synthetic records and start an audit. (2) The system converts observed attributes into short profiles and embeds both targets and auxiliary records. (3) ChromaDB retrieves the k closest candidates by cosine similarity. (4) Gemini receives the target and candidate evidence, then returns a candidate identity or abstains in a constrained JSON response [21], [22].')
replace(20,'(5) The verifier rejects invented IDs, checks field agreement, and applies score and margin thresholds. (6) JSON reports save the evidence, scores, flags, and decisions. (7) The optional anonymization module transforms records and audits them again against the same index. The evaluated changes are age bands, postal prefixes, medical or occupation suppression, and their combination. A global k-anonymity policy is an experimental baseline. Diagram examples of other sensitive fields are not benchmark variables. We measure information loss alongside linkage outcomes.')
replace(22,'The attacker knows age, postal code, occupation, gender, and medical flag; IDs are excluded from retrieval features. Knowing the medical flag is a strong assumption: the task is identity linkage, not diagnosis inference. When computing the release’s k-anonymity level, we exclude the medical flag, which the attacker is assumed to know. Local embeddings hash tokens and word pairs into 256 dimensions; cloud runs use gemini-embedding-001 and gemini-3.5-flash-lite at temperature zero [23]. The experiments use exact cosine search; the deployed application uses an approximate ChromaDB index.')
replace(23,'For an observed field f with value v, n(f,v) counts compatible index values and N is the index size. Its weight is w(f)=1+min{c, ln[(N+1)/(n(f,v)+1)]}. Let U and V be the total weights of agreeing and conflicting fields, and m the number observed. The evidence score is S=max{0,(U−V)/(U+V)}√(m/5), or zero when no usable evidence exists. Age intervals and postal prefixes allow compatible values; occupation paraphrases are still checked literally. Scores are rounded to six decimals, and remaining ties are broken by identity number.')
replace(24,'A candidate must meet score threshold τ and exceed the strongest retrieved competitor by margin γ. Gemini abstention is respected. Invalid JSON, unknown IDs, and provider failures produce no identity. Model confidence is not verified evidence. The application’s separate margin diagnostic min{1,max(0,Δ)/0.08} does not multiply S. These scores express support, not identification probabilities. If a cloud call fails, the experiment records an error instead of falling back to the local method.')
replace(39,'Each seed has 200 present and 200 absent validation identities, and 300 of each for testing. Present targets copy a sampled indexed identity before corruption. Absent identities are generated independently; accidental attribute matches are retained. All six variants of an identity remain in one split. Across five seeds, 2,000 validation identities yield 12,000 variants and 3,000 test identities yield 18,000. Clean copies provide a basic check, not evidence of robustness. The cloud cohorts remain unchanged.')
configs=[json.loads((RESULTS/str(s)/'configuration.json').read_text()) for s in SEEDS]
keys=sorted({c['local']['key'] for c in configs});taus=sorted({c['local']['tau'] for c in configs});gammas=sorted({c['local']['gamma'] for c in configs})
replace(45,'The always-top ablation disables abstention. The no-margin ablation removes the margin check and reselects the score threshold using validation data. The larger local cohorts reduce reliance on the earlier 20 absent validation identities. Selected settings are saved for each seed. For a fairer comparison, we also sweep each method’s test-score threshold while holding its validation-selected margin fixed (Fig. 3). These descriptive sweeps do not change the settings used in Table II.')
replace(47,'We count correct links for present targets (C), wrong links for present targets (W), present-target abstentions (P_A), false links for absent targets (F), and absent-target abstentions (A_A). Recall is C/(C+W+P_A), absent FPR is F/(F+A_A), and precision is C/(C+W+F). Coverage is the share of all targets assigned an identity. Ambiguous, insufficient-evidence, and error outcomes remain separate in saved predictions. Percentages are means ± sample SD across five populations; variants of one identity are not independent replications.')
replace(48,'Cloud experiments retain the original small cohorts selected before prediction: two present and two absent base identities per split and seed, each with six variants. This gives 24 validation and 24 test variants per seed. Gemini nomination and its hash control share candidates from the full 10,000-person index. The separate semantic pilot uses a smaller index. Quotas limited these experiments, so they do not establish semantic scalability. The benchmark prompt is standardized but differs from the dashboard prompt.')
replace(50,'The research package saves generated populations, target identities, split membership, corruptions, retrieved candidates, scores, decisions, and raw cloud responses. Each prediction row identifies its target, method, and policy. Counts and figures are calculated from these rows. Source hashes and model settings identify the evaluated code. Prompts exclude truth labels, target-group labels, and synthetic names.')
replace(51,'Seven evaluation checks cover disjoint identity splits, repeatable generation, agreement with application scores at cap three, complete precision denominators, failure handling, minimum k, and utility calculations for coarse intervals. Nineteen existing retrieval and anonymization checks also pass. The runner uses a separate collection. Cached model outputs preserve the evidence; new calls may differ in wording and latency.')
replace(54,'TABLE II. MAIN TEST RESULTS OVER FIVE POPULATIONS (%)')
names=['Exact match','Fellegi–Sunter','Hash + verification','No margin (ablation)','Always top candidate (ablation)']
for row,name,m in zip(doc.tables[2].rows[1:],names,['exact','fs','local','no_margin','top1']):
    for c,v in zip(row.cells,[name,ms(MAIN[m]['recall']),ms(MAIN[m]['fpr']),ms(MAIN[m]['precision'])]):cell(c,v)
replace(56,'Table II combines the six conditions within each seed, then summarizes variation across seeds. Its methods operate at different test false-link rates; recall alone cannot rank them fairly. Figure 3 compares their trade-offs. The always-top rule assigns an identity to every absent target, giving 100% absent FPR by design.')
replace(57,'TABLE III. COMPLETE OUTCOME COUNTS ON 18,000 VARIANTS')
for row,name,m in zip(doc.tables[3].rows[1:],names,['exact','fs','local','no_margin','top1']):
    for c,v in zip(row.cells,[name]+[MAIN[m]['totals'][k] for k in ['C','W','P_A','F','A_A']]):cell(c,v)
replace(59,f'Each row contains 9,000 present and 9,000 absent variants. Local verification issues {counts["C"]+counts["W"]+counts["F"]:,} identities: {counts["C"]:,} correct links, {counts["W"]:,} wrong present links, and {counts["F"]:,} absent false links. Its pooled precision is {100*counts["precision"]:.1f}%; Table II averages seed-level precision. At test FPR limits of 2% and 5%, mean recall is {fmt(MATCH["local",.02])}% and {fmt(MATCH["local",.05])}% for the local verifier, versus {fmt(MATCH["fs",.02])}% and {fmt(MATCH["fs",.05])}% for FS. FS has higher recall at both limits. These test sweeps do not select Table II settings.')
replace(61,'Fig. 2. Correct-link recall by condition: local verification (blue) and FS (orange). Whiskers show seed SD; each condition has 1,500 present and 1,500 absent variants.')
replace(62,'Clean-copy performance does not predict performance on changed records. Missing or altered occupations reduce support, while literal field checks can reject retrieved paraphrases. Saved candidate lists distinguish retrieval failures from verification failures. The no-margin and always-top ablations show the cost of issuing links with weaker evidence.')
replace(64,'Fig. 3. Recall versus absent FPR for both methods as score thresholds vary. Curves show mean best recall under each FPR limit with validation-selected margins fixed. Stars mark the Table II operating points; test sweeps are descriptive.')
replace(75,'Each original pilot index has 150 candidates: 50 present validation/test identities and 100 random distractors. Hash and semantic retrieval both reach 73.3% recall; adding Gemini gives 68.3% recall and 1.7% absent FPR. This small pilot shows no semantic advantage. Only Table V methods share the same candidate pool; comparison with Table II would mix index-size effects with method effects.')
replace(77,'We transform 600 clean test records per seed and re-audit them against the unchanged index. Policies use ten-year age bands, three-digit postal prefixes, medical or occupation suppression, and their combination. A global k=5 baseline follows a fixed ladder of broader age and postal values and categorical suppression until every released class has at least five rows. Tables VI–VII retain clean-validation thresholds. For Table VIII, both attackers reselect thresholds and margins on validation targets transformed by the same policy, using the same 5% validation error limits. The local score grid spans 0–1 and includes an abstain-all setting, allowing adjustment to missing fields.')
replace(78,'TABLE VI. FIXED-ATTACKER RESULTS AND INFORMATION LOSS')
for row,name,p in zip(doc.tables[6].rows[1:],NAMES,POLICIES):
    loss=statistics.mean(float(r['normalized_loss']) for r in UTILITY[p]);min_k=min(int(r['min_k']) for r in UTILITY[p])
    for c,v in zip(row.cells,[name,ms(MIT[p]['recall']),ms(MIT[p]['fpr']),f'{loss:.3f}',min_k]):cell(c,v)
kmin=min(int(r['min_k']) for r in UTILITY['k5']);kmax=max(int(r['min_k']) for r in UTILITY['k5'])
replace(80,f'Information loss averages five field penalties: zero for exact retention, one for suppression, age-interval width/67, and removed postal digits/6. Combined mitigation changes 80% of cells, suppresses 40%, and has loss 0.527, but minimum demographic k remains one. The global policy reaches minimum k of {kmin}–{kmax}, at greater information loss. These k values describe the released cohort.')
tv={p:statistics.mean(float(r['marginal_tv']) for r in UTILITY[p]) for p in POLICIES}
replace(81,f'We also compare age-decade and occupation counts. Intervals crossing decades remain unresolved rather than receiving invented exact ages. Mean total-variation error is {tv["age_band"]:.2f} for age bands, {tv["combined"]:.2f} for combined mitigation, and {tv["k5"]:.2f} for global k. This limited workload does not measure predictive utility.')
replace(82,'TABLE VII. FIXED-ATTACKER COUNTS ON 3,000 CLEAN RECORDS')
for row,name,p in zip(doc.tables[7].rows[1:],NAMES,POLICIES):
    for c,v in zip(row.cells,[name]+[MIT[p]['totals'][k] for k in ['C','W','P_A','F','A_A']]):cell(c,v)
replace(84,'Each policy has 1,500 present and 1,500 absent identities. Table VII counts issued and withheld identities; precision is undefined if no identity is issued. Table VIII reports mean recall/FPR (%) after policy-specific validation tuning. Both methods use the same transformed targets; FS searches the full index.')
# Insert a compact adaptive comparison after the existing outcome explanation.
caption=deepcopy(ps[82]._p);ps[84]._p.addnext(caption)
from docx.text.paragraph import Paragraph
cp=Paragraph(caption,ps[84]._parent);cp.clear();cp.add_run('TABLE VIII. ADAPTIVE RE-AUDIT RECALL / FPR (%)')
template=deepcopy(doc.tables[6]._tbl);caption.addnext(template)
from docx.table import Table
table=Table(template,doc._body)
for row in table.rows:
    for c in [c for c in row._tr if c.tag==qn('w:tc')][3:]:
        # Remove the last two cells, preserving row properties.
        if c.tag==qn('w:tc'):row._tr.remove(c)
for row in table.rows:
    for c in row.cells:
        c.width=Inches(1.12)
        for p in c.paragraphs:p.paragraph_format.space_after=Pt(0)
for c,v in zip(table.rows[0].cells,['Policy','Local R / FPR','FS R / FPR']):cell(c,v)
for row,name,p in zip(table.rows[1:],NAMES,POLICIES):
    vals=[name]+[f'{fmt(ADAPT[p][m]["recall"])} / {fmt(ADAPT[p][m]["fpr"])}' for m in ['adaptive_local','adaptive_fs']]
    for c,v in zip(row.cells,vals):cell(c,v)
# Keep its three-column grid consistent with the cells.
grid=table._tbl.tblGrid
for ch in list(grid):grid.remove(ch)
from docx.oxml import OxmlElement
for _ in range(3):
    ch=OxmlElement('w:gridCol');ch.set(qn('w:w'),'1620');grid.append(ch)

# Updated discussion uses measured values rather than the superseded small sample.
evidence=[]
for seed in SEEDS:
    raw=json.loads((RESULTS/str(seed)/'raw_mitigations.json').read_text());key=configs[SEEDS.index(seed)]['local']['key']
    evidence.append({p:statistics.mean(r['raw'][key]['score'] for r in raw if r['condition']==p) for p in ['original','age_band']})
es={p:statistics.mean(r[p] for r in evidence) for p in ['original','age_band']}
replace(87,f'Age bands broaden agreement: age 30–39 matches every candidate in that interval. Here, mean best-candidate evidence changes from {es["original"]:.3f} to {es["age_band"]:.3f}, while fixed-attacker recall changes from {fmt(MIT["original"]["recall"])}% to {fmt(MIT["age_band"]["recall"])}%. Ages 35 and 37 disagree exactly but agree after banding, so several identities may receive stronger support. Retrieval and frequency weights can change too. Wrong attribution is a potential harm, not automatically a privacy benefit.')
ac=ADAPT['combined'];fixedtaus=sorted({c['mitigation']['tau'] for c in configs})
replace(88,f'With combined mitigation, only three fields remain. The score cannot exceed √(3/5) ≈ 0.775. At the fixed threshold τ=0.80 in Tables VI–VII, the engine cannot issue a match on combined records. Zero recall is not proof of protection. Policy-specific validation tuning gives combined recall/FPR of {fmt(ac["adaptive_local"]["recall"])}%/{fmt(ac["adaptive_local"]["fpr"])}% for the local verifier and {fmt(ac["adaptive_fs"]["recall"])}%/{fmt(ac["adaptive_fs"]["fpr"])}% for FS. Minimum demographic k is still one. Excluding the known medical flag from k gives no sensitive-attribute guarantee.')
replace(89,'Synthetic populations and fixed corruption rules limit how far the findings generalize. Larger local cohorts improve the comparison, but cloud experiments still contain few independent identities. Model outputs can vary at temperature zero, and quotas limit semantic tests. Exact cosine search also differs from the deployed approximate index. Correlated variants do not establish statistical significance. Real-world benchmarks, repeated model calls, broader utility tests, and approximate-index evaluation are still needed.')
replace(91,'The engine brings together candidate retrieval, LLM assessment, field verification, abstention, and repeatable mitigation audits. Matched false-link comparisons and adaptive re-audits make its limits clearer. The main contribution is a workflow whose decisions can be inspected, rather than a claim that an LLM improves linkage accuracy. The small cloud study found equal recall with and without Gemini. Better semantic checks and larger cloud studies need further testing.')
replace(100,'[8] M. Ko, J. Jeong, S. S. Thakur, G. Kim, and R. Jia, “From weak cues to real identities: Evaluating inference-driven de-anonymization in LLM agents,” Proc. 43rd ICML, PMLR, vol. 306, pp. 59572–59589, 2026. https://proceedings.mlr.press/v306/ko26b.html')
replace(101,'[9] S. Lermen, D. Paleka, J. Swanson, M. Aerni, N. Carlini, and F. Tramèr, “Large-scale online deanonymization with LLMs,” USENIX Security, pp. 1947–1966, 2026; arXiv:2602.16800.')
for row in doc.tables[0].rows:
    for c in row.cells:
        for p in c.paragraphs:
            if p.text=='Jai Vinita':
                for r in p.runs:r.text=r.text.replace('Jai Vinita','Jai Vinita L*')
            if p.text=='jaivinita.l@vit.ac.in':
                p.add_run('\n*Corresponding author').font.size=Pt(8)

def font(size,bold=False):return ImageFont.truetype(str(Path('C:/Windows/Fonts')/('timesbd.ttf' if bold else 'times.ttf')),size)
def noise():
    conds=['clean','age_noise','missing','changed','combined_noise','paraphrase'];labels=['Clean','Age ±1','Missing occ.','Changed occ.','Combined','Paraphrase']
    im=Image.new('RGB',(1800,1020),'white');g=ImageDraw.Draw(im);left,right=380,1700
    for i,(cond,label) in enumerate(zip(conds,labels)):
        y=90+i*135;g.text((15,y+20),label,font=font(60),fill='black')
        for offset,method,color in [(0,'local','#165e83'),(45,'fs','#b84d22')]:
            vs=stats('predictions',method,cond)['recall'];x=left+statistics.mean(vs)*(right-left);sd=statistics.stdev(vs)*(right-left)
            g.rectangle((left,y+offset,x,y+offset+32),fill=color);g.line((max(left,x-sd),y+offset+16,min(right,x+sd),y+offset+16),fill='black',width=3)
    for v in [0,25,50,75,100]:g.text((left+v/100*(right-left)-30,905),str(v),font=font(52),fill='black')
    g.text((590,965),'Correct-link recall (%)',font=font(54),fill='black');p=ROOT/'noise.png';im.save(p);return p
def curves():
    im=Image.new('RGB',(1800,930),'white');g=ImageDraw.Draw(im);l,r,t,b=180,1730,75,740
    for y in [0,25,50,75,100]:
        yy=b-y/100*(b-t);g.line((l,yy,r,yy),fill='#dddddd',width=2);g.text((65,yy-25),str(y),font=font(50),fill='black')
    for x in [0,2,5,10,15,20]:
        xx=l+x/20*(r-l);g.text((xx-25,b+20),str(x),font=font(50),fill='black')
    for method,col in [('local','#165e83'),('fs','#b84d22')]:
        allrows={s:[x for x in read(s,'comparison_curve') if x['method']==method] for s in SEEDS};pts=[]
        for i in range(81):
            budget=i/400;ys=[max(float(x['recall']) for x in allrows[s] if float(x['fpr'])<=budget+1e-10) for s in SEEDS]
            pts.append((l+budget/.2*(r-l),b-statistics.mean(ys)*(b-t)))
        g.line(pts,fill=col,width=8)
        x=l+mean(MAIN[method]['fpr'])/20*(r-l);y=b-statistics.mean(MAIN[method]['recall'])*(b-t)
        star=[]
        for i in range(10):
            a=-math.pi/2+i*math.pi/5;rad=25 if i%2==0 else 11;star.append((x+rad*math.cos(a),y+rad*math.sin(a)))
        g.polygon(star,fill=col)
    g.text((15,10),'Recall (%)',font=font(50),fill='black');g.text((620,810),'Absent false-link rate (%)',font=font(55),fill='black')
    for x,col,label in [(220,'#165e83','Local verifier'),(1070,'#b84d22','Fellegi–Sunter')]:
        g.line((x,905,x+70,905),fill=col,width=8);g.text((x+90,875),label,font=font(50),fill='black')
    p=ROOT/'matched_curve.png';im.save(p);return p
for shape,path in zip(list(doc.inline_shapes)[1:],[noise(),curves()]):
    relid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[relid]._blob=path.read_bytes()

doc.save(OUTPUT)
report={'main':MAIN,'mitigation':MIT,'adaptive':ADAPT,'fixed_mitigation_thresholds':fixedtaus,'local_keys':keys,'local_taus':taus,'local_gammas':gammas,'matched':{f'{k[0]}_{k[1]}':v for k,v in MATCH.items()},'repository_url_pending':True}
(ROOT/'revision_numbers.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(OUTPUT)
print('Main',[(m,ms(MAIN[m]['recall']),ms(MAIN[m]['fpr'])) for m in MAIN])
print('Adaptive combined',[(m,fmt(ac[m]['recall']),fmt(ac[m]['fpr'])) for m in ac])
